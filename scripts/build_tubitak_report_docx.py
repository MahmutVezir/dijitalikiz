from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
BODY_PT = 12
DATA_TABLE_PT = 10
SCHEDULE_PT = 9
TITLE = "LLM Tabanlı Çapraz-Bileşen Hata Teşhisinde Yapılandırılmış Bağımlılık Temsilinin Etkisinin Deneysel İncelenmesi"
REQUIRED = [
    "Yapay Zeka Araçları Kullanım Beyanı", "Özet", "Anahtar Kelimeler",
    "Amaç", "Giriş", "Yöntem", "Proje İş-Zaman Çizelgesi", "Bulgular",
    "Sonuç ve Tartışma", "Öneriler", "Kaynaklar", "Ekler",
]
MONTHS = ["Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık", "Ocak"]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def normalize_label(text: str) -> str:
    value = unicodedata.normalize("NFKD", text.strip())
    return "".join(char for char in value if not unicodedata.combining(char)).replace("ı", "i").casefold()


def words(text: str) -> int:
    return len(re.findall(r"\b[\w’'-]+\b", text, re.UNICODE))


def set_font(run, size=BODY_PT, bold=None, italic=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for field in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn("w:" + field), name)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def inline(paragraph, text: str, size=BODY_PT):
    tick = chr(96)
    pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*|" + re.escape(tick) + r".*?" + re.escape(tick) + r")")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos:match.start()]), size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size, italic=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, size)
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]), size)


def cell_shading(cell, fill="D9EAF7"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def cell_borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn("w:" + edge))
        if elem is None:
            elem = OxmlElement("w:" + edge)
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "5")
        elem.set(qn("w:color"), "7F9DB9")


def set_cell(cell, value, size=DATA_TABLE_PT, bold=False, alignment=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if alignment is not None:
        p.alignment = alignment
    inline(p, value, size)
    for r in p.runs:
        if bold:
            r.bold = True
        set_font(r, size, bold=r.bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_data_table(doc, rows):
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for ri, row in enumerate(rows):
        vals = row + [""] * (width - len(row))
        for ci, val in enumerate(vals):
            cell = table.cell(ri, ci)
            align = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or ci > 0) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell(cell, val, DATA_TABLE_PT, bold=(ri == 0), alignment=align)
            cell_borders(cell)
            if ri == 0:
                cell_shading(cell)

    # Keep rows intact and repeat the column header on continuation pages.
    if len(table.rows) > 1:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if not tr_pr.xpath("./w:tblHeader"):
            header_flag = OxmlElement("w:tblHeader")
            header_flag.set(qn("w:val"), "true")
            tr_pr.append(header_flag)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if not tr_pr.xpath("./w:cantSplit"):
            tr_pr.append(OxmlElement("w:cantSplit"))

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)


def table_row(line):
    cells = [v.strip() for v in line.strip().strip("|").split("|")]
    return [v.replace("**", "").replace(chr(96), "").strip() for v in cells]


def is_separator(row):
    return bool(row) and all(re.fullmatch(r":?-{3,}:?", x.replace(" ", "")) for x in row)


def add_schedule_table(doc, schedule_xml, rows):
    if len(rows) < 2:
        raise ValueError("Schedule table needs at least one task.")
    if len(rows[0]) != 11 or normalize_label(rows[0][0]) != "isin adi":
        raise ValueError("Unexpected work-time table header.")
    if [normalize_label(x) for x in rows[0][1:]] != [normalize_label(x) for x in MONTHS]:
        raise ValueError("Schedule months do not match the official table template.")
    body = doc._element.body
    sectpr = body.sectPr
    body.insert(body.index(sectpr), deepcopy(schedule_xml))
    tbl = doc.tables[-1]
    tasks = rows[1:]
    while len(tbl.rows) < 2 + len(tasks):
        tbl._tbl.append(deepcopy(tbl.rows[-1]._tr))
    for ri, task in enumerate(tasks, start=2):
        vals = task + [""] * (11 - len(task))
        for ci in range(11):
            set_cell(tbl.cell(ri, ci), vals[ci], SCHEDULE_PT, alignment=WD_ALIGN_PARAGRAPH.CENTER if ci else WD_ALIGN_PARAGRAPH.LEFT)

    # Repeat the official two-row month header when a table must continue on
    # another page, and prevent any task row from splitting across pages.
    for ri in (0, 1):
        tr_pr = tbl.rows[ri]._tr.get_or_add_trPr()
        if not tr_pr.xpath("./w:tblHeader"):
            header_flag = OxmlElement("w:tblHeader")
            header_flag.set(qn("w:val"), "true")
            tr_pr.append(header_flag)
    for ri in range(2, len(tbl.rows)):
        tr_pr = tbl.rows[ri]._tr.get_or_add_trPr()
        if not tr_pr.xpath("./w:cantSplit"):
            tr_pr.append(OxmlElement("w:cantSplit"))

    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def page_number(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, 10)
    for kind, text in [("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)]:
        if text is not None:
            elem = OxmlElement("w:instrText" if text.strip() == "PAGE" else "w:t")
            if elem.tag == qn("w:instrText"):
                elem.set(qn("xml:space"), "preserve")
            elem.text = text
        else:
            elem = OxmlElement("w:fldChar")
            elem.set(qn("w:fldCharType"), kind)
        run._r.append(elem)


def setup(doc):
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)
        # Preserve the official template's embedded PAGE field. Inserting another
        # field into the secondary paragraph creates duplicate page numbers.
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    rpr = normal._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for field in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn("w:" + field), FONT)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(4)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "TÜBİTAK 2204-A proje raporu taslağı"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""


def extract_following_paragraph(lines, start):
    values = []
    i = start
    while i < len(lines):
        line = lines[i]
        if line.startswith("## ") or line.startswith("### "):
            break
        if line.strip():
            values.append(line.strip())
        elif values:
            break
        i += 1
    return " ".join(values)


def build_body(source, doc, schedule_xml):
    lines = source.splitlines()
    observed = [line[3:].strip() for line in lines if line.startswith("## ")]
    if observed != REQUIRED:
        raise ValueError(f"Report headings are not in official order: {observed}")
    if any("İçindekiler" in x for x in observed):
        raise ValueError("The report must not contain a contents section before the body.")

    abstract = extract_following_paragraph(lines, next(i for i, x in enumerate(lines) if x == "## Özet") + 1)
    keyword_text = extract_following_paragraph(lines, next(i for i, x in enumerate(lines) if x == "## Anahtar Kelimeler") + 1)
    metrics = {
        "abstract_word_count": words(abstract),
        "keyword_count": len([x for x in keyword_text.split(";") if x.strip()]),
        "project_title_word_count": words(TITLE),
        "section_order": observed,
    }
    if not 150 <= metrics["abstract_word_count"] <= 250:
        raise ValueError(f"Abstract word count outside 150-250: {metrics['abstract_word_count']}")
    if metrics["keyword_count"] > 5:
        raise ValueError(f"Keyword count exceeds five: {metrics['keyword_count']}")
    if metrics["project_title_word_count"] > 15:
        raise ValueError(f"Project title exceeds 15 words: {metrics['project_title_word_count']}")

    current_section = ""
    para_lines = []
    i = 0

    def flush():
        nonlocal para_lines
        if not para_lines:
            return
        value = " ".join(part.strip() for part in para_lines).strip()
        if value:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_together = True
            inline(p, value, BODY_PT)
        para_lines = []

    while i < len(lines):
        raw = lines[i]
        text = raw.strip()
        if not text:
            flush()
            i += 1
            continue
        if text.startswith("# "):
            i += 1
            continue
        if text.startswith(">"):
            flush()
            i += 1
            continue
        if text.startswith("## "):
            flush()
            current_section = text[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.0
            if current_section == "Proje İş-Zaman Çizelgesi":
                p.paragraph_format.page_break_before = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(current_section), BODY_PT, bold=True)
            i += 1
            continue
        if text.startswith("### "):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(text[4:].strip()), BODY_PT, bold=True)
            i += 1
            continue
        if text.startswith("|"):
            flush()
            raw_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw_rows.append(table_row(lines[i]))
                i += 1
            rows = [row for row in raw_rows if not is_separator(row)]
            if current_section == "Proje İş-Zaman Çizelgesi":
                add_schedule_table(doc, schedule_xml, rows)
            else:
                add_data_table(doc, rows)
            continue
        if re.match(r"^\s*[-*]\s+", raw):
            flush()
            content_text = re.sub(r"^\s*[-*]\s+", "", raw)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_font(p.add_run("• "), BODY_PT)
            inline(p, content_text, BODY_PT)
            i += 1
            continue
        ordered = re.match(r"^\s*(\d+)\.\s+(.*)$", raw)
        if ordered:
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_font(p.add_run(ordered.group(1) + ". "), BODY_PT)
            inline(p, ordered.group(2), BODY_PT)
            i += 1
            continue
        if text.startswith("**Proje Ana Alanı:**") or text.startswith("**Proje Tematik Alanı:**") or text.startswith("**Proje Adı (Başlığı):**"):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            inline(p, text, BODY_PT)
            i += 1
            continue
        para_lines.append(raw)
        if raw.endswith("  "):
            flush()
        i += 1
    flush()
    return metrics


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()

    source = args.markdown.read_text(encoding="utf-8")
    doc = Document(args.template)
    if not doc.tables:
        raise ValueError("Official template schedule table is missing.")
    schedule_xml = deepcopy(doc.tables[0]._tbl)
    body = doc._element.body
    sectpr = body.sectPr
    for child in list(body):
        if child is not sectpr:
            body.remove(child)
    setup(doc)
    metrics = build_body(source, doc, schedule_xml)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    check = Document(args.output)
    text = "\n".join(p.text for p in check.paragraphs)
    if "NOT: Kırmızı ile yazılı kısımlar" in text:
        raise ValueError("Template instruction text remains in the report.")
    if "Yapay Zeka Araçları Kullanım Beyanı" not in text:
        raise ValueError("AI-use declaration heading is missing.")
    if len(check.sections) != 1 or len(check.tables) != 4:
        raise ValueError(f"Unexpected document structure: sections={len(check.sections)}, tables={len(check.tables)}")
    footer_fields = check.sections[0].footer._element.xpath(".//w:instrText")
    page_fields = [item for item in footer_fields if "PAGE" in (item.text or "").upper()]
    if len(page_fields) != 1:
        raise ValueError(f"Expected one official page-number field, found {len(page_fields)}.")
    sec = check.sections[0]
    for value in (sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin):
        if abs(value.cm - 2.5) > 0.02:
            raise ValueError("All page margins must be 2.5 cm.")
    manifest = {
        "source_markdown_sha256": sha256_file(args.markdown),
        "official_template_sha256": sha256_file(args.template),
        "output_docx_sha256": sha256_file(args.output),
        "font": "Times New Roman",
        "body_font_size_pt": 12,
        "page_size": "A4",
        "margins_cm": {"top": round(sec.top_margin.cm, 3), "bottom": round(sec.bottom_margin.cm, 3), "left": round(sec.left_margin.cm, 3), "right": round(sec.right_margin.cm, 3)},
        "template_guidance_removed": True,
        "table_count": len(check.tables),
        **metrics,
    }
    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    args.manifest.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
